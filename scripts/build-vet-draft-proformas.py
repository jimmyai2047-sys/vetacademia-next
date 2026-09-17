# Builds 5 bilingual DRAFT field-veterinarian proformas as .docx (python-docx)
# and .pdf (reportlab) into vetacademia-next/public/proformas/.
# All drafts are MODEL formats (not official GOI/AHD forms) and are labelled so.
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, HRFlowable)
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = r"D:\VetAcademia (VA)\vetacademia-next\public\proformas"
os.makedirs(OUT, exist_ok=True)

TTC = r"C:\Windows\Fonts\Nirmala.ttc"
pdfmetrics.registerFont(TTFont("Nirmala", TTC, subfontIndex=0))
pdfmetrics.registerFont(TTFont("Nirmala-Bold", TTC, subfontIndex=1))

GREEN = RGBColor(0x00, 0x5F, 0x48)
DRAFT_BANNER = ("DRAFT — Model format for field use. Not an official GOI / State AHD form. "
                "Verify with your State Animal Husbandry Department before official use. / "
                "प्रारूप — केवल क्षेत्रीय उपयोग हेतु नमूना। आधिकारिक उपयोग से पूर्व राज्य पशुपालन विभाग से सत्यापित करें।")

# ---------------- content model ----------------
# Each doc: title_en, title_hi, intro, sections = list of ("lines"|"table"|"sign", payload)
DOCS = [
    {
        "base": "vaccination-certificate-animal",
        "title_en": "Vaccination Certificate — Individual Animal (Draft)",
        "title_hi": "टीकाकरण प्रमाण पत्र — व्यक्तिगत पशु (प्रारूप)",
        "src": ("Model format based on common AHD practice and AQCS export-documentation norms "
                "(animal ID, vaccine batch, next due date). Ref: dahd.gov.in; aqcsindia.gov.in."),
        "sections": [
            ("lines", ["Certificate No.: ____________      Date / दिनांक: ____________",
                        "Owner name / स्वामी का नाम: ________________________   Phone: ____________",
                        "Address / पता: ________________________________________________________"]),
            ("table", [["Ear Tag / कान टैग", "Species / प्रजाति", "Breed / नस्ल", "Age / आयु", "Sex / लिंग", "Colour & Marks / रंग-निशान"],
                        ["", "", "", "", "", ""]]),
            ("table", [["Disease / रोग", "Vaccine / टीका", "Batch No.", "Mfg / Exp", "Dose & Route", "Site"],
                        ["", "", "", "", "", ""],
                        ["", "", "", "", "", ""]]),
            ("lines", ["Date of vaccination / टीकाकरण दिनांक: ____________      Next due / अगला टीका: ____________",
                        "Cold chain maintained at 2–8 °C: Yes / No      Adverse reaction observed: Yes / No (details: ________)",
                        "Note: Brucella vaccine only in female calves aged 4–8 months. FMD booster every 6 months; PPR / HS / BQ as per state schedule."]),
            ("sign", None),
        ],
    },
    {
        "base": "fitness-certificate-domestic",
        "title_en": "Fitness Certificate for Domestic Movement / Market (Draft)",
        "title_hi": "स्वास्थ्य एवं योग्यता प्रमाण पत्र — आवागमन / बाज़ार (प्रारूप)",
        "src": ("Model format based on common AHD practice. For export, only the AQCS Official Veterinarian "
                "may sign the importing country's health certificate (Ref: aqcsindia.gov.in)."),
        "sections": [
            ("lines", ["Certificate No.: ____________      Examination date & place: ________________________",
                        "Purpose / उद्देश्य (tick):  Sale / Market-Mela / Transport / Show / Other ________",
                        "Owner name / स्वामी: ________________________   Phone: ____________"]),
            ("table", [["S.No.", "Ear Tag / ID", "Species", "Breed", "Age", "Sex", "Pregnancy status"],
                        ["1", "", "", "", "", "", ""],
                        ["2", "", "", "", "", "", ""],
                        ["3", "", "", "", "", "", ""]]),
            ("lines", ["Clinical findings (tick examined):  Temperature ___°F   Pulse ___/min   Respiration ___/min",
                        "Skin & coat: Normal / Abnormal ___   Eyes & discharge: Normal / Abnormal ___",
                        "Gait & limbs: Normal / Lameness ___   Appetite & rumination: Normal / Abnormal ___",
                        "Visible signs of infectious / contagious disease: None / Suspected (details: ________________)",
                        "Opinion:  FIT / UNFIT for the above purpose.   Valid for ______ days from examination.",
                        "Conditions: transport as per Transport of Animals Rules; isolate and re-examine if signs appear."]),
            ("sign", None),
        ],
    },
    {
        "base": "pregnancy-diagnosis-certificate",
        "title_en": "Pregnancy Diagnosis Certificate (Draft)",
        "title_hi": "गर्भावस्था निदान प्रमाण पत्र (प्रारूप)",
        "src": "Model format based on common AHD practice (AI records + per-rectal / USG finding).",
        "sections": [
            ("lines", ["Certificate No.: ____________      Examination date / दिनांक: ____________",
                        "Owner / स्वामी: ________________________   Phone: ____________",
                        "Animal: Species ________  Breed ________  Age ________  Ear Tag ________"]),
            ("table", [["Breeding date / प्रजनन तिथि", "Method (AI / Natural)", "Sire / Semen batch", "Inseminator"],
                        ["", "", "", ""],
                        ["", "", "", ""]]),
            ("lines", ["Examination method (tick):  Per-rectal / Ultrasonography",
                        "Finding (tick):  PREGNANT (~ ______ months)  /  NON-PREGNANT  /  DOUBTFUL — recheck on ____________",
                        "Expected calving date / संभावित ब्यांत तिथि: ____________",
                        "Advice: balanced feeding in last trimester; deworming and vaccination as per schedule."]),
            ("sign", None),
        ],
    },
    {
        "base": "outbreak-intimation-sample-checklist",
        "title_en": "Disease Outbreak Intimation & Sample Dispatch Checklist (Draft)",
        "title_hi": "पशुरोग प्रकोप सूचना एवं नमूना प्रेषण सूची (प्रारूप)",
        "src": ("Model field aid aligned to NADRS 2.0 workflow (Block VO: daily incidence + FIR on outbreak; "
                "escalate via DVO to State and CDRMU New Delhi) and the Prevention and Control of Infectious and "
                "Contagious Diseases in Animals Act, 2009. Ref: nadrs.dahd.gov.in; nivedi.res.in (NADRES)."),
        "sections": [
            ("lines", ["Intimation No.: ____________      Date: ____________      Reported by (name/designation/phone): ________________________",
                        "Village: ____________  Block: ____________  District: ____________  State: ____________",
                        "Farm / Gaushala / Poultry farm: ________________________  Owner: ________________________"]),
            ("table", [["Species", "Population at risk", "Sick", "Dead", "Date of onset"],
                        ["Cattle/Buffalo", "", "", "", ""],
                        ["Sheep/Goat", "", "", "", ""],
                        ["Poultry", "", "", "", ""],
                        ["Other ______", "", "", "", ""]]),
            ("lines", ["Clinical signs seen (tick):  Fever / Vesicles on mouth-foot / Abortions / Sudden death / Skin nodules / "
                        "Respiratory distress / Diarrhoea / Nervous signs / Drop in milk-egg / Other ____________",
                        "Provisional diagnosis: ________________________",
                        "Control measures taken (tick):  Isolation of sick / Movement stopped / Disinfection / Ring vaccination / "
                        "Treatment started / Carcass disposed (deep burial/incineration) / Neighbouring villages alerted"]),
            ("table", [["Sample (Serum plain vial / EDTA blood / Nasal-ocular swab / Tissue in 10% formalin / Post-mortem tissue chilled)", "Animal ID", "Packing (Chilled 2-8C / Formalin)", "Lab (RDDL / State lab / IVRI / NIVEDI)"],
                        ["", "", "", ""],
                        ["", "", "", ""],
                        ["", "", "", ""]]),
            ("lines", ["Packing rules: serum and swabs CHILLED (2-8 C, icebox); one tissue set in 10% formalin (histopathology); "
                        "label every vial (ID, date, village); dispatch within 24 hours with this form.",
                        "Reported to: Block VO ______ (phone ______)  DVO ______ (phone ______)  NADRS 2.0 portal entry done: Yes / No (FIR No. ______)",
                        "Zoonotic suspicion? Yes / No — if yes, inform Medical / IDSP District Surveillance Officer same day."]),
            ("sign", None),
        ],
    },
    {
        "base": "drug-withdrawal-quick-chart",
        "title_en": "Drug Withdrawal Period — Quick Reference Chart (Draft)",
        "title_hi": "औषधि प्रतीक्षा अवधि — संदर्भ तालिका (प्रारूप)",
        "src": ("Indicative chart only. Always follow the product label. FSSAI (Contaminants, Toxins and Residues) Regulations, 2011 "
                "as amended 2018 fix tolerance limits for 100+ veterinary drugs (e.g. milk limits as low as 0.01 mg/kg). "
                "BANNED in food animals: Chloramphenicol, Nitrofurans, Metronidazole, Clenbuterol, DES, Glycopeptides. "
                "Ref: fssai.gov.in; vetworld 2021 review of FSSAI MRLs."),
        "sections": [
            ("lines", ["READ FIRST: values below are conservative field guidance. The product LABEL is final. "
                        "Never send milk/meat of treated animals to market during withdrawal. / लेबल ही अंतिम है।"]),
            ("table", [["Drug (common field use)", "Milk", "Meat", "Remarks"],
                        ["Oxytetracycline LA", "4 days", "28 days", "Most common residue violator — observe strictly"],
                        ["Enrofloxacin", "4 days", "10 days", "Poultry meat 10 days"],
                        ["Amoxicillin", "3 days (72 h)", "14 days", "Verify label of combination products"],
                        ["Procaine penicillin G", "3 days", "10 days", "Allergic keep-out for milk"],
                        ["Ceftiofur", "Nil (most labels)", "0-4 days", "Confirm nil-milk on YOUR label"],
                        ["Gentamicin", "5 days", "60+ days", "Long tissue persistence — avoid in dairy"],
                        ["Ivermectin (cattle)", "Do NOT use in lactating dairy", "35-49 days", "Check label; long withdrawal"],
                        ["Albendazole", "3 days", "14 days", "Verify label"],
                        ["Meloxicam", "5 days", "15 days", "Verify label"],
                        ["Dexamethasone", "3 days (72 h)", "21 days", "Also: abortion risk in late pregnancy"],
                        ["Xylazine", "3 days", "3 days", "Verify label"]]),
            ("sign", None),
        ],
    },
]

SIGN_LINES = ["Examined / Certified by (Veterinarian / पशु चिकित्सक): ________________________",
              "Registration No. / पंजीकरण सं.: ____________   Seal / मुहर: ____________",
              "Signature / हस्ताक्षर: ____________   Date / दिनांक: ____________   Place: ____________"]

# ---------------- docx builder ----------------
def build_docx(d):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Nirmala UI"
    st.font.size = Pt(10.5)
    h = doc.add_heading(d["title_en"], level=1)
    for r in h.runs:
        r.font.color.rgb = GREEN
    p = doc.add_paragraph()
    r = p.add_run(d["title_hi"])
    r.bold = True
    r.font.size = Pt(13)
    ban = doc.add_paragraph()
    r = ban.add_run(DRAFT_BANNER)
    r.font.size = Pt(8)
    r.italic = True
    for kind, payload in d["sections"]:
        if kind == "lines":
            for line in payload:
                doc.add_paragraph(line)
        elif kind == "table":
            t = doc.add_table(rows=len(payload), cols=len(payload[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate(payload):
                for j, cell in enumerate(row):
                    c = t.cell(i, j)
                    c.text = ""
                    pr = c.paragraphs[0].add_run(cell)
                    pr.font.size = Pt(9.5)
                    if i == 0:
                        pr.bold = True
            doc.add_paragraph("")
        elif kind == "sign":
            doc.add_paragraph("")
            for s in SIGN_LINES:
                doc.add_paragraph(s)
    doc.add_paragraph("")
    src = doc.add_paragraph()
    r = src.add_run("Source / basis: " + d["src"])
    r.font.size = Pt(8)
    r.italic = True
    path = os.path.join(OUT, d["base"] + ".docx")
    doc.save(path)
    print("DOCX", path, os.path.getsize(path))

# ---------------- pdf builder ----------------
def build_pdf(d):
    path = os.path.join(OUT, d["base"] + ".pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=14 * mm, bottomMargin=14 * mm,
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            title=d["title_en"], author="VetAcademia (draft)")
    sTitle = ParagraphStyle("t", fontName="Nirmala-Bold", fontSize=14, textColor=colors.HexColor("#005f48"),
                            spaceAfter=2 * mm, leading=17)
    sHi = ParagraphStyle("h", fontName="Nirmala-Bold", fontSize=11, spaceAfter=2 * mm, leading=14)
    sBody = ParagraphStyle("b", fontName="Nirmala", fontSize=9, leading=12, spaceAfter=1.5 * mm)
    sSmall = ParagraphStyle("s", fontName="Nirmala", fontSize=7, leading=9, textColor=colors.grey,
                            spaceBefore=3 * mm)
    sCell = ParagraphStyle("c", fontName="Nirmala", fontSize=8, leading=10)
    sCellH = ParagraphStyle("ch", fontName="Nirmala-Bold", fontSize=8, leading=10,
                            textColor=colors.white)
    story = [Paragraph(d["title_en"], sTitle), Paragraph(d["title_hi"], sHi),
             HRFlowable(width="100%", thickness=1, color=colors.HexColor("#005f48")),
             Spacer(1, 2 * mm), Paragraph(DRAFT_BANNER, sSmall)]
    for kind, payload in d["sections"]:
        if kind == "lines":
            for line in payload:
                story.append(Paragraph(line.replace(" / ", " / "), sBody))
        elif kind == "table":
            head, rows = payload[0], payload[1:]
            data = [[Paragraph(c, sCellH) for c in head]]
            for row in rows:
                data.append([Paragraph(c or " ", sCell) for c in row])
            w = (170 * mm) / len(head)
            t = Table(data, colWidths=[w] * len(head), repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#005f48")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(Spacer(1, 2 * mm))
            story.append(t)
            story.append(Spacer(1, 2 * mm))
        elif kind == "sign":
            story.append(Spacer(1, 3 * mm))
            for s in SIGN_LINES:
                story.append(Paragraph(s, sBody))
    story.append(Paragraph("Source / basis: " + d["src"], sSmall))
    story.append(Paragraph("Generated by VetAcademia as a DRAFT model format.", sSmall))
    doc.build(story)
    print("PDF ", path, os.path.getsize(path))

for d in DOCS:
    build_docx(d)
    build_pdf(d)
print("DONE", len(DOCS), "drafts")
